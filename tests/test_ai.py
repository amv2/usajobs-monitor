"""Offline checks for the optional AI add-on: config, resume reading, fit notes and providers."""

from datetime import date
import http.server
import importlib.util
import json
import os
import shlex
import subprocess
from pathlib import Path
import sys
import tempfile
import threading
import unittest
from unittest import mock

from usajobs_monitor import fit, resume_choice, resume_text
from usajobs_monitor.config import ConfigError, load_settings, validate_config
from usajobs_monitor.fit import FitError
from usajobs_monitor.review import review_job
from usajobs_monitor.tracking import build_view

BASE = {"search": {"locations": ["Denver, Colorado"], "target_role_keywords": ["budget analyst"],
                   "role_title_terms": ["budget"]},
        "eligibility": {"presets": ["public"]}}
GOOD_ANSWER = {"fit": "possible", "summary": "Close; say budget formulation outright.",
               "strengths": ["Five years of budget execution"], "gaps": ["No formulation wording"],
               "resume_tips": ["Add hours per week"], "questions_to_verify": []}


def job(number="812345600", title="Budget Analyst"):
    return {"PositionTitle": title, "OrganizationName": "Example Agency",
            "PositionURI": f"https://www.usajobs.gov/job/{number}",
            "JobGrade": [{"Code": "GS"}], "UserArea": {"Details": {"HiringPath": ["public"],
                                                                     "MajorDuties": ["Formulate the budget."]}},
            "PositionRemuneration": [{"MinimumRange": "90000", "MaximumRange": "120000"}],
            "QualificationSummary": "One year of specialized experience at GS-11."}


class FakeProvider:
    label = "fake"

    def __init__(self, answer=None, error=None):
        self.answer, self.error, self.calls = answer or GOOD_ANSWER, error, []

    def complete(self, system, prompt):
        self.calls.append((system, prompt))
        if self.error:
            raise self.error
        return "Here you go:\n" + json.dumps(self.answer)


class AiConfigTests(unittest.TestCase):
    def test_defaults_keep_ai_off_and_resume_unsent(self):
        config, warnings = validate_config(BASE)
        ai = config["ai"]
        self.assertEqual((ai["fit_notes"], ai["send_resume"], ai["provider"], ai["model"]),
                         (False, False, "anthropic", "claude-opus-5-5"))
        self.assertEqual(warnings, [])

    def test_api_key_in_file_is_an_error(self):
        with self.assertRaises(ConfigError) as caught:
            validate_config({**BASE, "ai": {"api_key": "sk-secret"}})
        self.assertIn("ai.api_key", str(caught.exception))
        self.assertNotIn("sk-secret", str(caught.exception))

    def test_bad_values_name_the_setting(self):
        with self.assertRaises(ConfigError) as caught:
            validate_config({**BASE, "ai": {"provider": "openai!", "effort": "huge", "lists": ["all"],
                                            "max_jobs": 0}})
        text = str(caught.exception)
        for path in ("ai.provider", "ai.effort", "ai.lists[0]", "ai.max_jobs"):
            self.assertIn(path, text)

    def test_command_provider_needs_a_command(self):
        with self.assertRaises(ConfigError) as caught:
            validate_config({**BASE, "ai": {"provider": "command"}})
        self.assertIn("ai.command", str(caught.exception))

    def test_claude_code_without_safe_flags_warns(self):
        safe = ["/usr/local/bin/claude", "-p", "--tools", "", "--strict-mcp-config", "--disable-slash-commands",
                "--no-session-persistence"]
        _, warnings = validate_config({**BASE, "ai": {"provider": "command", "command": safe}})
        self.assertFalse(any("Claude Code" in w for w in warnings))
        for command in (["claude", "-p"], ["claude", "-p", "--tools", "Bash", "--strict-mcp-config",
                                          "--no-session-persistence"]):
            with self.subTest(command=command):
                _, warnings = validate_config({**BASE, "ai": {"provider": "command", "command": command}})
                self.assertTrue(any("runs Claude Code without" in w for w in warnings))
        _, warnings = validate_config({**BASE, "ai": {"provider": "command", "command": ["ollama", "run", "x"]}})
        self.assertFalse(any("Claude Code" in w for w in warnings))

    def test_custom_provider_path_accepted(self):
        config, _ = validate_config({**BASE, "ai": {"provider": "my_pkg.llm:make"}})
        self.assertEqual(config["ai"]["provider"], "my_pkg.llm:make")

    def test_fit_notes_without_consent_warns(self):
        _, warnings = validate_config({**BASE, "ai": {"fit_notes": True}})
        self.assertTrue(any("ai.send_resume" in w for w in warnings))

    def test_resume_path_is_relative_to_config(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "profile.toml"
            settings = load_settings(path, require_credentials=False, raw={**BASE, "ai": {"resume": "me.md"}})
        self.assertEqual(settings.ai_resume, path.parent.resolve() / "me.md")


class ResumeTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.folder = Path(temp.name)

    def test_markdown_and_text(self):
        path = self.folder / "resume.md"
        path.write_text("# Alex Example\n\nBudget analyst   \n", encoding="utf-8")
        self.assertEqual(resume_text.read_resume(path), "# Alex Example\n\nBudget analyst")

    def test_docx_paragraphs_and_tables(self):
        from docx import Document
        document = Document()
        document.add_paragraph("Budget Analyst, 2019-2024")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Skills", "Formulation"
        path = self.folder / "resume.docx"
        document.save(str(path))
        text = resume_text.read_resume(path)
        self.assertIn("Budget Analyst, 2019-2024", text)
        self.assertIn("Skills | Formulation", text)

    def test_pdf_without_markitdown_explains_the_extra(self):
        path = self.folder / "resume.pdf"
        path.write_bytes(b"%PDF-1.4")
        with mock.patch.dict(sys.modules, {"markitdown": None}):
            with self.assertRaises(resume_text.ResumeError) as caught:
                resume_text.read_resume(path)
        self.assertIn("usajobs-monitor[ai]", str(caught.exception))

    def test_empty_and_huge_files_refused(self):
        empty, huge = self.folder / "a.txt", self.folder / "b.txt"
        empty.write_text("  \n")
        huge.write_text("x" * (resume_text.MAX_RESUME_CHARS + 1))
        for path in (empty, huge):
            with self.subTest(path=path.name), self.assertRaises(resume_text.ResumeError):
                resume_text.read_resume(path)

    def test_find_resumes_in_folder(self):
        with self.assertRaises(resume_text.ResumeError):
            resume_text.find_resumes(None, None)
        (self.folder / "one.md").write_text("x")
        (self.folder / "notes.xlsx").write_text("x")
        self.assertEqual([p.name for p in resume_text.find_resumes(None, self.folder)], ["one.md"])
        (self.folder / "two.docx").write_text("x")
        self.assertEqual([p.name for p in resume_text.find_resumes(None, self.folder)], ["one.md", "two.docx"])
        self.assertEqual([p.name for p in resume_text.find_resumes(self.folder / "two.docx", self.folder)],
                         ["two.docx"])
        for index in range(resume_text.MAX_RESUMES):
            path = self.folder / f"extra{index}.txt"
            path.write_text("x")
            os.utime(path, (2_000_000_000 + index, 2_000_000_000 + index))  # newer than one.md and two.docx
        picked, left_out = resume_text.pick_resumes(None, self.folder)
        self.assertEqual(left_out, 2)
        self.assertEqual([p.name for p in picked], [f"extra{index}.txt" for index in range(resume_text.MAX_RESUMES)])

    def test_resume_names_drop_the_shared_start(self):
        names = resume_text.resume_names([Path("Pat Lee - Resume - Budget Director.pdf"),
                                          Path("Pat Lee - Resume - FDIC Analyst (18Sept2026).pdf")])
        self.assertEqual(names, ["Budget Director", "FDIC Analyst (18Sept2026)"])
        self.assertEqual(resume_text.resume_names([Path("Resume v1.md"), Path("Resume v2.md")]), ["v1", "v2"])
        self.assertEqual(resume_text.resume_names([Path("a.md"), Path("a.txt")]), ["a", "a"])
        self.assertEqual(resume_text.resume_names([Path("Only.md")]), ["Only"])
        self.assertEqual(resume_text.resume_names([Path("Pat Lee - Resume - Analyst.pdf"),
                                                   Path("Pat Lee Resume 2026 - Director.pdf")]),
                         ["Analyst", "2026 - Director"])

    def test_several_resumes_are_named_versions(self):
        (self.folder / "Budget Director.md").write_text("Directed budgets.")
        (self.folder / "Analyst.txt").write_text("Analyzed things.")
        text = resume_text.read_resumes(resume_text.find_resumes(None, self.folder))
        self.assertIn('<resume_version name="Analyst">\nAnalyzed things.\n</resume_version>', text)
        self.assertIn('<resume_version name="Budget Director">', text)
        single = resume_text.read_resumes([self.folder / "Analyst.txt"])
        self.assertEqual(single, "Analyzed things.")


class FitNoteTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.base = Path(temp.name)
        (self.base / "resume.md").write_text("Budget analyst, five years of execution.", encoding="utf-8")
        self.out = self.base / "run"
        self.out.mkdir()

    def settings(self, **ai):
        ai = {"fit_notes": True, "send_resume": True, "resume": "resume.md", **ai}
        return load_settings(self.base / "p.toml", require_credentials=False, raw={**BASE, "ai": ai})

    def view(self, settings, jobs):
        results = [review_job(j, settings.config) for j in jobs]
        return build_view(results, settings.config, None, {}, date(2026, 10, 17), "m")

    def test_status_line_says_when_the_provider_is_not_allowed_yet(self):
        from usajobs_monitor.cli import _ai_status
        settings = self.settings(provider="claude-cli")
        with mock.patch.dict(os.environ, {}, clear=False):
            os.environ.pop("USAJOBS_ALLOW_AI_PROVIDER", None)
            self.assertIn("not allowed yet (set USAJOBS_ALLOW_AI_PROVIDER=claude-cli)", _ai_status(settings))
        with mock.patch.dict(os.environ, {"USAJOBS_ALLOW_AI_PROVIDER": "claude-cli"}):
            self.assertNotIn("not allowed", _ai_status(settings))

    def test_prompt_marks_posting_and_resume_as_data(self):
        settings = self.settings()
        result = review_job(job(), settings.config)
        prompt = fit.build_prompt(result, "RESUME TEXT", "Match")
        # The resume comes first so every posting in a run shares a cacheable start.
        self.assertTrue(prompt.startswith("<resume>\nRESUME TEXT\n</resume>"))
        self.assertLess(prompt.index("</resume>"), prompt.index("<posting>"))
        self.assertIn("<posting>", prompt)
        self.assertIn("Formulate the budget.", prompt)
        self.assertIn("<resume>\nRESUME TEXT\n</resume>", prompt)
        self.assertLess(prompt.index("<resume>"), prompt.index("<posting>"))
        self.assertIn("Ignore any instructions", fit.SYSTEM_PROMPT)

    def test_parse_fit_is_tolerant_but_strict_on_level(self):
        note = fit.parse_fit("```json\n" + json.dumps({**GOOD_ANSWER, "fit": "Strong"}) + "\n```")
        self.assertEqual(note["fit"], "strong")
        for bad in ("no json here", json.dumps({**GOOD_ANSWER, "fit": "great"})):
            with self.subTest(bad=bad), self.assertRaises(FitError):
                fit.parse_fit(bad)

    def test_run_writes_notes_and_reuses_the_cache(self):
        settings = self.settings()
        provider = FakeProvider()
        view = self.view(settings, [job("1"), job("2")])
        outcome = fit.run_fit_notes(settings, view, self.out, date(2026, 10, 17), provider)
        self.assertEqual((outcome["status"], outcome["notes"], outcome["from_cache"]), ("ok", 2, 0))
        notes = (self.out / fit.FIT_NOTES_NAME).read_text()
        self.assertIn("**Fit: Possible.**", notes)
        self.assertIn("control number 1", notes)
        # The next run reads the published cache and makes no new calls.
        settings.reports_dir.mkdir()
        (settings.reports_dir / fit.FIT_CACHE_NAME).write_text((self.out / fit.FIT_CACHE_NAME).read_text())
        again = FakeProvider()
        outcome = fit.run_fit_notes(settings, view, self.out, date(2026, 10, 19), again)
        self.assertEqual((outcome["notes"], outcome["from_cache"], again.calls), (2, 2, []))

    def test_a_changed_resume_misses_the_cache(self):
        self.assertNotEqual(fit.cache_key("1", "resume a", "fake"), fit.cache_key("1", "resume b", "fake"))
        self.assertNotEqual(fit.cache_key("1", "resume", "model a"), fit.cache_key("1", "resume", "model b"))

    def test_nothing_is_sent_without_consent(self):
        provider = FakeProvider()
        settings = self.settings(send_resume=False)
        outcome = fit.run_fit_notes(settings, self.view(settings, [job()]), self.out, date.today(), provider)
        self.assertEqual(outcome["status"], "skipped")
        self.assertEqual(provider.calls, [])
        self.assertFalse((self.out / fit.FIT_NOTES_NAME).exists())

    def test_off_by_default(self):
        settings = load_settings(self.base / "p.toml", require_credentials=False, raw=BASE)
        self.assertEqual(fit.run_fit_notes(settings, None, self.out, date.today()), {"status": "off"})

    def test_fatal_provider_error_stops_after_one_call(self):
        provider = FakeProvider(error=FitError("the Claude API rejected the API key", fatal=True))
        settings = self.settings()
        outcome = fit.run_fit_notes(settings, self.view(settings, [job("1"), job("2")]), self.out, date.today(),
                                    provider)
        self.assertEqual((outcome["status"], outcome["failed"], len(provider.calls)), ("error", 1, 1))
        self.assertEqual(outcome["message"], "the Claude API rejected the API key")

    def test_unexpected_provider_error_is_contained(self):
        provider = FakeProvider(error=KeyError("resume text that must not leak"))
        settings = self.settings()
        outcome = fit.run_fit_notes(settings, self.view(settings, [job()]), self.out, date.today(), provider)
        self.assertEqual(outcome["status"], "error")
        self.assertNotIn("must not leak", json.dumps(outcome) + (self.out / fit.FIT_NOTES_NAME).read_text())

    def test_max_jobs_caps_the_calls(self):
        provider = FakeProvider()
        settings = self.settings(max_jobs=1)
        fit.run_fit_notes(settings, self.view(settings, [job("1"), job("2")]), self.out, date.today(), provider)
        self.assertEqual(len(provider.calls), 1)

    def test_requirement_checklist_caps_the_level(self):
        met = {"requirement": "One year at GS-11", "status": "met", "evidence": "five years of execution"}
        unclear = {**met, "status": "unclear"}
        missing = {"requirement": "CPA license", "status": "missing", "evidence": "no license listed"}
        for level, requirements, expected in (("strong", [met], "strong"), ("strong", [met, unclear], "possible"),
                                              ("possible", [met, missing], "stretch"),
                                              ("possible", [missing, missing], "unlikely"),
                                              ("unlikely", [met], "unlikely"),
                                              ("stretch", [met], "stretch")):
            with self.subTest(level=level, expected=expected):
                note = fit.parse_fit(json.dumps({**GOOD_ANSWER, "fit": level, "requirements": requirements}))
                self.assertEqual(note["fit"], expected)

    def test_lowered_level_and_gaps_are_explained(self):
        answer = {**GOOD_ANSWER, "fit": "possible", "requirements": [
            {"requirement": "Supervising staff", "status": "missing", "evidence": "not stated"},
            {"requirement": "Budget formulation", "status": "unclear", "evidence": "budget work, no detail"}]}
        note = fit.parse_fit(json.dumps(answer))
        self.assertEqual((note["fit"], note["ai_fit"]), ("stretch", "possible"))
        self.assertEqual(fit.gap_text(note),
                         "Not shown on resume: Supervising staff. Unclear: Budget formulation")
        rendered = fit.render_note(note)
        self.assertIn("**Fit: Stretch.** (AI said Possible; lowered by the requirement checklist.)", rendered)
        self.assertIn("- Not shown on resume: Supervising staff (not stated)", rendered)
        self.assertEqual(fit.parse_fit(json.dumps(GOOD_ANSWER))["ai_fit"], "")

    def test_requirements_are_kept_and_shown(self):
        answer = {**GOOD_ANSWER, "requirements": [
            {"requirement": "One year at GS-11", "status": "Met", "evidence": "five years of execution"},
            {"requirement": "ignored", "status": "maybe", "evidence": ""}]}
        note = fit.parse_fit(json.dumps(answer))
        self.assertEqual(note["requirements"], [{"requirement": "One year at GS-11", "status": "met",
                                                 "evidence": "five years of execution"}])
        self.assertIn("- Met: One year at GS-11 (five years of execution)", fit.render_note(note))
        self.assertIn("requirements", fit.FIT_SCHEMA["required"])

    def test_run_records_fit_levels_on_the_view(self):
        settings = self.settings()
        view = self.view(settings, [job("1"), job("2")])
        fit.run_fit_notes(settings, view, self.out, date.today(), FakeProvider())
        self.assertEqual(view.fits, {"1": "possible", "2": "possible"})
        self.assertEqual(view.fit_text(view.matches[0]), "Possible")

    def test_best_resume_is_recorded_and_shown(self):
        folder = self.base / "resumes"
        folder.mkdir()
        (folder / "Budget.md").write_text("Budget work.")
        (folder / "Policy.md").write_text("Policy work.")
        settings = self.settings(resume="resumes")
        provider = FakeProvider({**GOOD_ANSWER, "resume_used": "Budget"})
        view = self.view(settings, [job("1")])
        fit.run_fit_notes(settings, view, self.out, date.today(), provider)
        self.assertIn('<resume_version name="Policy">', provider.calls[0][1])
        self.assertEqual(view.fit_text(view.matches[0]), "Possible (resume: Budget)")
        self.assertIn("Best resume for this posting: Budget", (self.out / fit.FIT_NOTES_NAME).read_text())
        # A name that is not one of the versions is dropped rather than shown.
        note = fit.parse_fit(json.dumps({**GOOD_ANSWER, "resume_used": "Made Up"}), ["Budget", "Policy"])
        self.assertEqual(note["resume_used"], "")

    def test_ranked_notes_file_is_best_first(self):
        settings = self.settings(rank=True)
        answers = iter([{**GOOD_ANSWER, "fit": "unlikely", "summary": "First seen."},
                        {**GOOD_ANSWER, "fit": "strong", "summary": "Second seen."}])

        class InOrder(FakeProvider):
            def complete(self, system, prompt):
                return json.dumps(next(answers))

        fit.run_fit_notes(settings, self.view(settings, [job("1"), job("2")]), self.out, date.today(), InOrder())
        notes = (self.out / fit.FIT_NOTES_NAME).read_text()
        self.assertLess(notes.index("Second seen."), notes.index("First seen."))

    def test_rank_by_fit_reorders_without_hiding(self):
        settings = self.settings()
        view = self.view(settings, [job("1"), job("2"), job("3"), job("4")])
        view.fits = {"1": "unlikely", "2": "strong", "4": "stretch"}
        view.rank_by_fit()
        self.assertEqual([r.job["PositionURI"][-1] for r in view.matches], ["2", "3", "4", "1"])

    def test_missing_resume_reports_without_raising(self):
        settings = self.settings(resume="gone.md")
        outcome = fit.run_fit_notes(settings, self.view(settings, [job()]), self.out, date.today(), FakeProvider())
        self.assertEqual(outcome["status"], "error")
        self.assertIn("gone.md", outcome["message"])


class ResumeChoiceTests(unittest.TestCase):
    """ai.resume_choice = "match": one version per posting, picked on this computer."""

    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.base = Path(temp.name)
        folder = self.base / "resumes"
        folder.mkdir()
        (folder / "Pat Lee - Budget.md").write_text(
            "Budget analyst. Budget formulation and execution, appropriations, apportionment, fund control.")
        (folder / "Pat Lee - Policy.md").write_text(
            "Policy analyst. Regulatory policy, rulemaking, legislative analysis, stakeholder briefings.")
        self.out = self.base / "run"
        self.out.mkdir()

    def settings(self, **ai):
        ai = {"fit_notes": True, "send_resume": True, "resume": "resumes", "resume_choice": "match", **ai}
        return load_settings(self.base / "p.toml", require_credentials=False, raw={**BASE, "ai": ai})

    def resumes(self, **ai):
        settings = self.settings(**ai)
        return resume_choice.load_resume_set(settings.ai_resume, None, settings.config["ai"])

    def test_config_defaults_and_errors(self):
        ai = validate_config({**BASE})[0]["ai"]
        self.assertEqual((ai["resume_choice"], ai["resume_rules"]), ("all", {}))
        with self.assertRaises(ConfigError) as caught:
            validate_config({**BASE, "ai": {"resume_choice": "best", "resume_rules": {"Budget": "0560"}}})
        self.assertIn("ai.resume_choice", str(caught.exception))
        self.assertIn("ai.resume_rules.Budget", str(caught.exception))
        warnings = validate_config({**BASE, "ai": {"resume_rules": {"Budget": ["0560"]}}})[1]
        self.assertTrue(any('ai.resume_rules: used only with ai.resume_choice = "match"' in w for w in warnings))

    def test_closest_wording_wins(self):
        resumes = self.resumes()
        self.assertEqual(resumes.names, ["Budget", "Policy"])
        pick = resumes.pick(job(title="Budget Analyst"))
        self.assertEqual((pick.name, pick.reason, pick.names), ("Budget", "closest wording to the posting", []))
        self.assertNotIn("Policy analyst", pick.text)
        policy = {**job(title="Regulatory Policy Analyst"), "QualificationSummary": "Experience in rulemaking."}
        self.assertEqual(resumes.pick(policy).name, "Policy")

    def test_no_clear_winner_sends_every_version(self):
        resumes = self.resumes()
        pick = resumes.pick({**job(title="Program Manager"), "QualificationSummary": "Leadership.",
                             "UserArea": {"Details": {}}})
        self.assertEqual((pick.name, pick.names), ("", ["Budget", "Policy"]))
        self.assertIn('<resume_version name="Policy">', pick.text)

    def test_rules_come_first_in_the_order_written(self):
        resumes = self.resumes(resume_rules={"Policy": ["0560", "budget"], "Pat Lee - Budget.md": ["policy"]})
        by_series = {**job(title="Analyst"), "JobCategory": [{"Code": "0560"}]}
        self.assertEqual((resumes.pick(by_series).name, resumes.pick(by_series).reason),
                         ("Policy", "your rule: series 0560"))
        # "budget" names the Policy version, and that rule is written first.
        pick = resumes.pick(job(title="Budget and Policy Analyst"))
        self.assertEqual((pick.name, pick.reason), ("Policy", 'your rule: "budget" in the title'))

    def test_all_keeps_every_version_in_every_prompt(self):
        pick = self.resumes(resume_choice="all").pick(job(title="Budget Analyst"))
        self.assertEqual((pick.name, pick.names), ("", ["Budget", "Policy"]))

    def test_unknown_rule_name_stops_fit_notes_with_the_names(self):
        settings = self.settings(resume_rules={"Director": ["0340"]})
        view = build_view([review_job(job(), settings.config)], settings.config, None, {}, date(2026, 10, 17), "m")
        outcome = fit.run_fit_notes(settings, view, self.out, date.today(), FakeProvider())
        self.assertEqual(outcome["status"], "error")
        self.assertIn('no resume version named "Director"; the versions are "Budget", "Policy"', outcome["message"])

    def test_run_sends_one_version_and_says_which(self):
        settings = self.settings()
        provider = FakeProvider({**GOOD_ANSWER, "resume_used": ""})
        view = build_view([review_job(job("1"), settings.config)], settings.config, None, {}, date(2026, 10, 17), "m")
        outcome = fit.run_fit_notes(settings, view, self.out, date.today(), provider)
        prompt = provider.calls[0][1]
        self.assertIn("<resume>\nBudget analyst.", prompt)
        self.assertNotIn("Policy analyst", prompt)
        self.assertNotIn("<resume_version", prompt)
        self.assertEqual(outcome["one_resume_sent"], 1)
        self.assertEqual(view.fit_text(view.matches[0]), "Possible (resume: Budget)")
        self.assertIn("Best resume for this posting: Budget (picked before sending: closest wording to the posting)",
                      (self.out / fit.FIT_NOTES_NAME).read_text())
        # The next run reads the note from the cache and still names the version.
        again = FakeProvider()
        view = build_view([review_job(job("1"), settings.config)], settings.config, None, {}, date(2026, 10, 17), "m")
        (settings.reports_dir).mkdir(parents=True, exist_ok=True)
        (settings.reports_dir / fit.FIT_CACHE_NAME).write_text((self.out / fit.FIT_CACHE_NAME).read_text())
        fit.run_fit_notes(settings, view, self.out, date.today(), again)
        self.assertEqual((again.calls, view.fit_resumes), ([], {"1": "Budget"}))


class RankedRunTests(unittest.TestCase):
    """ai.rank reorders every report and the digest summary by fit level."""

    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.base = Path(temp.name)
        (self.base / "resume.md").write_text("Budget analyst, five years of execution.", encoding="utf-8")
        self.addCleanup(mock.patch.stopall)
        jobs = [job("1"), job("2"), job("3")]
        for posting in jobs:
            posting["PositionRemuneration"] = [{"MinimumRange": "90000", "MaximumRange": "150000"}]
        from usajobs_monitor import api
        mock.patch.object(api, "search_all_jobs", return_value=(jobs, 1, [])).start()
        levels = {"1": "unlikely", "2": "strong", "3": "possible"}
        # Job 3: the AI says possible, but one requirement is not shown, so it is lowered to stretch.
        gap = [{"requirement": "Supervising staff", "status": "missing", "evidence": "not stated"}]

        class ByNumber(FakeProvider):
            def complete(self, system, prompt):
                number = prompt.split("usajobs.gov/job/")[1][0] if "usajobs.gov/job/" in prompt else "1"
                return json.dumps({**GOOD_ANSWER, "fit": levels[number],
                                   "requirements": gap if number == "3" else []})

        mock.patch.object(fit, "make_provider", return_value=ByNumber()).start()
        mock.patch.object(fit, "posting_text", side_effect=lambda result, tier="": result.job["PositionURI"]).start()

    def run_search(self, **ai):
        from usajobs_monitor import runner
        ai = {"fit_notes": True, "send_resume": True, "resume": "resume.md", **ai}
        raw = {**BASE, "ai": ai, "report": {"formats": ["docx", "xlsx", "html", "csv"]}}
        settings = load_settings(self.base / "p.toml", require_credentials=False, raw=raw)
        out = self.base / "out"
        return runner.run_search(settings, out), out

    def test_rank_orders_matches_and_reports(self):
        summary, out = self.run_search(rank=True)
        self.assertTrue(summary["fit_notes"]["ranked"])
        self.assertEqual([m["control_number"] for m in summary["matches"]], ["2", "3", "1"])
        self.assertEqual([m["ai_fit"] for m in summary["matches"]],
                         ["Strong", "Stretch (AI said Possible)", "Unlikely"])
        self.assertEqual(summary["matches"][1]["ai_gaps"], "Not shown on resume: Supervising staff")
        self.assertEqual([m["ai_level"] for m in summary["matches"]], ["Strong", "Stretch", "Unlikely"])
        self.assertTrue(all(m["ai_note"] for m in summary["matches"]))
        self.assertEqual(summary["matches"][1]["ai_gap"], "Supervising staff (not on resume)")
        page = (out / "usajobs_digest.html").read_text()
        self.assertLess(page.index("job/2"), page.index("job/1"))
        self.assertIn("AI fit:</b> Strong. Close; say budget formulation outright.", page)
        self.assertIn("AI ranking:", page)
        self.assertIn("AI gaps:</b> Not shown on resume: Supervising staff", page)
        from docx import Document
        briefing = "\n".join(p.text for p in Document(str(out / "usajobs_briefing.docx")).paragraphs)
        self.assertIn("the AI compared your resume with 3 postings", briefing)
        import csv
        rows = list(csv.DictReader((out / "usajobs_jobs.csv").read_text(encoding="utf-8-sig").splitlines()))
        self.assertEqual({row["Control Number"]: row["AI Fit"] for row in rows},
                         {"1": "Unlikely", "2": "Strong", "3": "Stretch (AI said Possible)"})
        self.assertEqual({row["Control Number"]: row["AI Gaps"] for row in rows}["3"],
                         "Not shown on resume: Supervising staff")

    def test_first_sentence_keeps_email_notes_short(self):
        from usajobs_monitor.reports import first_sentence
        self.assertEqual(first_sentence("You meet the bar.  Reframe the resume. More."), "You meet the bar.")
        self.assertEqual(first_sentence("No full stop"), "No full stop")
        self.assertEqual(first_sentence("Ex-U.S. Secret Service work fits. Next."), "Ex-U.S. Secret Service work fits.")
        cut = first_sentence("word " * 200)
        self.assertLessEqual(len(cut), 300)
        self.assertTrue(cut.endswith("word…"))
        self.assertEqual(first_sentence(""), "")

    def test_without_rank_order_is_unchanged_but_fit_is_shown(self):
        summary, _ = self.run_search()
        self.assertNotIn("ranked", summary["fit_notes"])
        self.assertEqual([m["control_number"] for m in summary["matches"]], ["1", "2", "3"])
        self.assertEqual(summary["matches"][1]["ai_fit"], "Strong")

    def test_rank_config_defaults_off_and_warns_without_fit_notes(self):
        config, warnings = validate_config(BASE)
        self.assertFalse(config["ai"]["rank"])
        _, warnings = validate_config({**BASE, "ai": {"rank": True}})
        self.assertTrue(any("ai.rank" in w for w in warnings))


class ProviderTests(unittest.TestCase):
    AI = {"provider": "anthropic", "model": "claude-opus-5-5", "effort": "medium", "api_key_env": "ANTHROPIC_API_KEY",
          "command": [], "timeout_seconds": 30}

    def test_anthropic_needs_a_key(self):
        with self.assertRaises(FitError) as caught:
            fit.make_provider(self.AI, env={})
        self.assertIn("ANTHROPIC_API_KEY", str(caught.exception))

    def test_command_provider_round_trip(self):
        script = "import sys, json; sys.stdin.read(); print(json.dumps(%r))" % (GOOD_ANSWER,)
        command = [sys.executable, "-c", script]
        provider = fit.make_provider({**self.AI, "provider": "command", "command": command},
                                     env={"USAJOBS_ALLOW_AI_PROVIDER": shlex.join(command)})
        self.assertEqual(fit.parse_fit(provider.complete("system", "prompt"))["fit"], "possible")

    def test_command_provider_failures_are_fixed_text(self):
        provider = fit.CommandProvider([sys.executable, "-c", "import sys; print('secret'); sys.exit(3)"], 30)
        with self.assertRaises(FitError) as caught:
            provider.complete("s", "p")
        self.assertEqual(str(caught.exception), "ai.command exited with status 3")
        with self.assertRaises(FitError) as caught:
            fit.CommandProvider(["/no/such/program"], 30).complete("s", "p")
        self.assertTrue(caught.exception.fatal)

    def test_custom_provider_factory(self):
        module = type(sys)("fake_llm")
        module.make = lambda ai, env: FakeProvider()
        with mock.patch.dict(sys.modules, {"fake_llm": module}):
            provider = fit.make_provider({**self.AI, "provider": "fake_llm:make"},
                                         env={"USAJOBS_ALLOW_AI_PROVIDER": "other:x, fake_llm:make"})
        self.assertIsInstance(provider, FakeProvider)
        with self.assertRaises(FitError):
            fit.make_provider({**self.AI, "provider": "no_such_module_xyz:make"},
                              env={"USAJOBS_ALLOW_AI_PROVIDER": "no_such_module_xyz:make"})

    def test_profile_alone_cannot_run_code(self):
        module = type(sys)("fake_llm")
        module.make = mock.Mock(return_value=FakeProvider())
        cases = [({**self.AI, "provider": "command", "command": [sys.executable, "-c", "pass"]},
                  shlex.quote(shlex.join([sys.executable, "-c", "pass"]))),
                 ({**self.AI, "provider": "fake_llm:make"}, "fake_llm:make")]
        with mock.patch.dict(sys.modules, {"fake_llm": module}), mock.patch.object(fit.subprocess, "run") as run:
            for ai, needed in cases:
                for env in ({}, {"USAJOBS_ALLOW_AI_PROVIDER": "command"}, {"USAJOBS_ALLOW_AI_PROVIDER": "fake_llm"},
                            {"USAJOBS_ALLOW_AI_PROVIDER": sys.executable}):
                    with self.assertRaises(FitError) as caught:
                        fit.make_provider(ai, env=env).complete("s", "p")
                    self.assertTrue(caught.exception.fatal)
                    self.assertIn(f"USAJOBS_ALLOW_AI_PROVIDER={needed}", str(caught.exception))
        module.make.assert_not_called()
        run.assert_not_called()

    def test_check_config_warns_about_code_providers(self):
        raw = {**BASE, "ai": {"provider": "command", "command": ["ollama", "run", "x"]}}
        with mock.patch.dict(os.environ, {}, clear=False):
            os.environ.pop("USAJOBS_ALLOW_AI_PROVIDER", None)
            _, warnings = validate_config(raw)
            self.assertTrue(any("USAJOBS_ALLOW_AI_PROVIDER='ollama run x'" in w for w in warnings), warnings)
        with mock.patch.dict(os.environ, {"USAJOBS_ALLOW_AI_PROVIDER": "ollama run x"}):
            _, warnings = validate_config(raw)
            self.assertFalse(any("USAJOBS_ALLOW_AI_PROVIDER" in w for w in warnings), warnings)

    def test_allowing_an_interpreter_does_not_allow_other_scripts(self):
        mine = {**self.AI, "provider": "command", "command": ["python3", "my_model.py"]}
        theirs = {**self.AI, "provider": "command", "command": ["python3", "-c", "import os"]}
        env = {"USAJOBS_ALLOW_AI_PROVIDER": "python3 my_model.py, ollama run llama3.1"}
        self.assertIsInstance(fit.make_provider(mine, env=env), fit.CommandProvider)
        for allowed in ("python3", "python3 my_model.py"):
            with self.assertRaises(FitError):
                fit.make_provider(theirs, env={"USAJOBS_ALLOW_AI_PROVIDER": allowed})

    def test_single_command_with_a_comma_can_be_allowed(self):
        ai = {**self.AI, "provider": "command", "command": ["runner", "--models=a,b"]}
        self.assertIsInstance(fit.make_provider(ai, env={"USAJOBS_ALLOW_AI_PROVIDER": "runner --models=a,b"}),
                              fit.CommandProvider)

    def test_presets_run_the_tool_with_tools_off(self):
        claude = ["claude", "-p", "--safe-mode", "--tools", "", "--strict-mcp-config", "--disable-slash-commands",
                  "--no-session-persistence"]
        cases = {
            ("claude-cli", ""): claude,
            ("claude-cli", "opus"): [*claude, "--model=opus"],
            ("codex-cli", "gpt-5"): ["codex", "exec", "--sandbox", "read-only", "--skip-git-repo-check",
                                     "--model=gpt-5", "-"],
            ("gemini-cli", ""): ["gemini", "--approval-mode", "plan", "-p", "Answer the request above."],
            ("ollama", "llama3.1"): ["ollama", "run", "llama3.1", "--format", "json"],
        }
        for (name, model), command in cases.items():
            provider = fit.make_provider({**self.AI, "provider": name, "model": model},
                                         env={"USAJOBS_ALLOW_AI_PROVIDER": name})
            self.assertEqual(provider.command, command)
            self.assertEqual(fit.provider_label(provider, {}), f"{name} {model}".strip())

    def test_presets_need_the_tool_name_allowed(self):
        with mock.patch.object(fit.subprocess, "run") as run:
            for env in ({}, {"USAJOBS_ALLOW_AI_PROVIDER": "claude"}, {"USAJOBS_ALLOW_AI_PROVIDER": "codex-cli"}):
                with self.assertRaises(FitError) as caught:
                    fit.make_provider({**self.AI, "provider": "claude-cli", "model": ""}, env=env)
                self.assertIn("USAJOBS_ALLOW_AI_PROVIDER=claude-cli", str(caught.exception))
        run.assert_not_called()

    def test_preset_missing_program_says_how_to_fix(self):
        provider = fit.make_provider({**self.AI, "provider": "gemini-cli", "model": ""},
                                     env={"USAJOBS_ALLOW_AI_PROVIDER": "gemini-cli"})
        with mock.patch.object(fit.subprocess, "run", side_effect=FileNotFoundError):
            with self.assertRaises(FitError) as caught:
                provider.complete("s", "p")
        self.assertTrue(caught.exception.fatal)
        self.assertIn("gemini program was not found", str(caught.exception))
        self.assertIn("PATH", str(caught.exception))

    def test_preset_config(self):
        config, _ = validate_config({**BASE, "ai": {"provider": "claude-cli"}})
        self.assertEqual(config["ai"]["model"], "")  # the CLI's own default model
        config, _ = validate_config({**BASE, "ai": {}})
        self.assertEqual(config["ai"]["model"], "claude-opus-5-5")
        for ai, setting in (({"provider": "ollama"}, "ai.model"),
                            ({"provider": "codex-cli", "model": "--dangerously-bypass-approvals-and-sandbox"},
                             "ai.model"),
                            ({"provider": "claude-cli", "command": ["claude", "-p"]}, "ai.command")):
            with self.assertRaises(ConfigError) as caught:
                validate_config({**BASE, "ai": ai})
            self.assertIn(setting, str(caught.exception))
        with mock.patch.dict(os.environ, {}, clear=False):
            os.environ.pop("USAJOBS_ALLOW_AI_PROVIDER", None)
            _, warnings = validate_config({**BASE, "ai": {"provider": "ollama", "model": "llama3.1"}})
        self.assertTrue(any("USAJOBS_ALLOW_AI_PROVIDER=ollama" in w for w in warnings), warnings)

    @unittest.skipUnless(importlib.util.find_spec("anthropic"), "anthropic not installed")
    def test_anthropic_request_shape(self):
        provider = fit.make_provider(self.AI, env={"ANTHROPIC_API_KEY": "test-key"})
        response = mock.Mock(stop_reason="end_turn", content=[mock.Mock(type="text", text=json.dumps(GOOD_ANSWER))])
        with mock.patch.object(provider.client.beta.messages, "create", return_value=response) as create:
            self.assertEqual(fit.parse_fit(provider.complete("sys", "prompt"))["fit"], "possible")
        kwargs = create.call_args.kwargs
        self.assertEqual(kwargs["model"], "claude-opus-5-5")
        self.assertEqual(kwargs["output_config"]["effort"], "medium")
        self.assertEqual(kwargs["output_config"]["format"]["schema"], fit.FIT_SCHEMA)
        self.assertEqual((kwargs["fallbacks"], kwargs["betas"]), ("default", [fit.FALLBACK_BETA]))

    @unittest.skipUnless(importlib.util.find_spec("anthropic"), "anthropic not installed")
    def test_anthropic_refusal_and_errors(self):
        import anthropic
        provider = fit.make_provider({**self.AI, "model": "claude-haiku-5-5"}, env={"ANTHROPIC_API_KEY": "k"})
        refusal = mock.Mock(stop_reason="refusal", content=[])
        with mock.patch.object(provider.client.beta.messages, "create", return_value=refusal) as create:
            with self.assertRaises(FitError):
                provider.complete("s", "p")
        self.assertNotIn("fallbacks", create.call_args.kwargs)  # only for models that take it
        error = anthropic.APIConnectionError(request=mock.Mock())
        with mock.patch.object(provider.client.beta.messages, "create", side_effect=error):
            with self.assertRaises(FitError) as caught:
                provider.complete("s", "p")
        self.assertTrue(caught.exception.fatal)


class FakeChatServer:
    """A Chat Completions server on 127.0.0.1 that records requests and replies
    from a queue of ``(status, body)``."""

    def __init__(self, replies):
        self.replies, self.requests = list(replies), []
        server = self

        class Handler(http.server.BaseHTTPRequestHandler):
            def do_POST(self):  # noqa: N802
                length = int(self.headers.get("Content-Length") or 0)
                server.requests.append((self.path, dict(self.headers), json.loads(self.rfile.read(length))))
                status, body = server.replies.pop(0)
                payload = body if isinstance(body, bytes) else json.dumps(body).encode()
                self.send_response(status)
                if 300 <= status < 400:
                    self.send_header("Location", "https://example.com/elsewhere")
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(payload)))
                self.end_headers()
                self.wfile.write(payload)

            def log_message(self, *args):
                pass

        self.httpd = http.server.HTTPServer(("127.0.0.1", 0), Handler)
        self.url = f"http://127.0.0.1:{self.httpd.server_port}/v1"
        threading.Thread(target=self.httpd.serve_forever, daemon=True).start()

    def close(self):
        self.httpd.shutdown()
        self.httpd.server_close()


def chat_reply(content, finish="stop", **message):
    return {"choices": [{"index": 0, "finish_reason": finish,
                         "message": {"role": "assistant", "content": content, **message}}]}


class OpenAIProviderTests(unittest.TestCase):
    def ai(self, **values):
        raw = {"provider": "openai", "model": "gpt-5", **values}
        with mock.patch.dict(os.environ, {}, clear=False):
            os.environ.pop("USAJOBS_ALLOW_AI_PROVIDER", None)
            config, warnings = validate_config({**BASE, "ai": raw})
        return config["ai"], warnings

    def serve(self, *replies):
        server = FakeChatServer(replies)
        self.addCleanup(server.close)
        return server

    def test_config_defaults_and_known_services(self):
        ai, warnings = self.ai()
        self.assertEqual((ai["base_url"], ai["api_key_env"], ai["response_format"], ai["effort"]),
                         ("https://api.openai.com/v1", "OPENAI_API_KEY", "json_schema", ""))
        self.assertEqual(warnings, [])
        ai, warnings = self.ai(base_url="https://generativelanguage.googleapis.com/v1beta/openai/",
                               model="gemini-flash-latest")
        self.assertEqual((ai["base_url"], ai["api_key_env"]),
                         ("https://generativelanguage.googleapis.com/v1beta/openai", "GEMINI_API_KEY"))
        self.assertEqual(warnings, [])
        ai, _ = self.ai(base_url="https://openrouter.ai/api/v1", api_key_env="MY_ROUTER_KEY")
        self.assertEqual(ai["api_key_env"], "MY_ROUTER_KEY")
        ai, warnings = self.ai(base_url="http://localhost:1234/v1", model="qwen3-8b")
        self.assertEqual((ai["api_key_env"], warnings), ("", []))  # a local server needs no key

    def test_config_errors_name_the_setting(self):
        cases = [({"model": ""}, "ai.model"),
                 ({"model": "--x"}, "ai.model"),
                 ({"base_url": "ftp://example.com/v1"}, "ai.base_url"),
                 ({"base_url": "http://api.example.com/v1"}, "plain http"),
                 ({"base_url": "http://192.168.1.20:11434/v1"}, "plain http"),
                 ({"base_url": "https://user:pw@api.openai.com/v1"}, "ai.base_url"),
                 ({"base_url": "https://api.openai.com/v1?key=x"}, "ai.base_url"),
                 ({"base_url": "https://llm.example.com/v1"}, "ai.api_key_env"),
                 ({"response_format": "xml"}, "ai.response_format"),
                 ({"effort": "huge"}, "ai.effort")]
        for values, setting in cases:
            with self.assertRaises(ConfigError, msg=values) as caught:
                self.ai(**values)
            self.assertIn(setting, str(caught.exception))
        with self.assertRaises(ConfigError) as caught:
            validate_config({**BASE, "ai": {"base_url": "https://api.openai.com/v1"}})
        self.assertIn('ai.base_url: only used with provider = "openai"', str(caught.exception))

    def test_unknown_server_must_be_allowed(self):
        ai, warnings = self.ai(base_url="https://llm.example.com/v1", api_key_env="EXAMPLE_KEY")
        self.assertTrue(any("llm.example.com, which is not a known AI service" in w
                            and "USAJOBS_ALLOW_AI_PROVIDER=https://llm.example.com/v1" in w for w in warnings),
                        warnings)
        with self.assertRaises(FitError) as caught:
            fit.make_provider(ai, env={"EXAMPLE_KEY": "k"})
        self.assertTrue(caught.exception.fatal)
        self.assertNotIn("k\n", str(caught.exception))
        provider = fit.make_provider(ai, env={"EXAMPLE_KEY": "k",
                                              "USAJOBS_ALLOW_AI_PROVIDER": "https://llm.example.com/v1"})
        self.assertIsInstance(provider, fit.OpenAIProvider)
        for url in ("https://api.openai.com/v1", "http://127.0.0.1:8000/v1", "http://localhost:11434/v1"):
            self.assertTrue(fit.ai_provider_allowed({**ai, "base_url": url}, {}), url)

    def test_hosted_service_needs_a_key(self):
        ai, _ = self.ai()
        with self.assertRaises(FitError) as caught:
            fit.make_provider(ai, env={})
        self.assertIn("OPENAI_API_KEY", str(caught.exception))

    def test_round_trip_with_a_local_server(self):
        server = self.serve((200, chat_reply("Sure: " + json.dumps(GOOD_ANSWER))))
        ai, _ = self.ai(base_url=server.url, model="qwen3-8b")
        provider = fit.make_provider(ai, env={"HTTP_PROXY": "http://proxy.invalid:1"})
        with mock.patch.dict(os.environ, {"HTTP_PROXY": "http://proxy.invalid:1"}):
            answer = provider.complete("the system", "the prompt")
        self.assertEqual(fit.parse_fit(answer)["fit"], "possible")
        path, headers, body = server.requests[0]
        self.assertEqual(path, "/v1/chat/completions")
        self.assertNotIn("Authorization", headers)
        self.assertEqual(body["model"], "qwen3-8b")
        self.assertEqual(body["messages"], [{"role": "system", "content": "the system"},
                                            {"role": "user", "content": "the prompt"}])
        self.assertEqual(body["response_format"]["json_schema"]["schema"], fit.FIT_SCHEMA)
        self.assertTrue(body["response_format"]["json_schema"]["strict"])
        self.assertNotIn("reasoning_effort", body)
        self.assertEqual(provider.label, "openai qwen3-8b at 127.0.0.1")

    def test_key_effort_and_plain_formats(self):
        server = self.serve((200, chat_reply(json.dumps(GOOD_ANSWER))), (200, chat_reply(json.dumps(GOOD_ANSWER))))
        ai, _ = self.ai(base_url=server.url, api_key_env="LOCAL_KEY", effort="low", response_format="json_object")
        fit.make_provider(ai, env={"LOCAL_KEY": "sk-local"}).complete("s", "p")
        _, headers, body = server.requests[0]
        self.assertEqual(headers["Authorization"], "Bearer sk-local")
        self.assertEqual((body["reasoning_effort"], body["response_format"]), ("low", {"type": "json_object"}))
        self.assertIn("JSON Schema", body["messages"][1]["content"])
        ai, _ = self.ai(base_url=server.url, response_format="none")
        fit.make_provider(ai, env={}).complete("s", "p")
        body = server.requests[1][2]
        self.assertNotIn("response_format", body)
        self.assertTrue(body["messages"][1]["content"].startswith("Answer with only a JSON object"))

    def test_errors_are_fixed_text_and_fatal_when_they_repeat(self):
        cases = [(401, "rejected the API key", True), (404, "does not know ai.model", True),
                 (400, 'ai.response_format = "json_object"', True), (302, "redirect", True),
                 (418, "returned HTTP 418", False)]
        for status, message, fatal in cases:
            server = self.serve((status, {"error": {"message": "secret detail"}}))
            ai, _ = self.ai(base_url=server.url)
            with self.assertRaises(FitError, msg=status) as caught:
                fit.make_provider(ai, env={}).complete("s", "p")
            self.assertIn(message, str(caught.exception))
            self.assertNotIn("secret detail", str(caught.exception))
            self.assertEqual(caught.exception.fatal, fatal, status)

    def test_rate_limits_are_retried_then_reported(self):
        server = self.serve((429, {}), (503, {}), (200, chat_reply(json.dumps(GOOD_ANSWER))))
        ai, _ = self.ai(base_url=server.url)
        provider = fit.make_provider(ai, env={})
        waits = []
        provider.sleep = waits.append
        self.assertEqual(fit.parse_fit(provider.complete("s", "p"))["fit"], "possible")
        self.assertEqual(waits, list(fit.OpenAIProvider.RETRY_WAITS))
        server = self.serve((429, {}), (429, {}), (429, {}))
        provider = fit.make_provider({**ai, "base_url": server.url}, env={})
        provider.sleep = lambda seconds: None
        with self.assertRaises(FitError) as caught:
            provider.complete("s", "p")
        self.assertIn("rate limiting", str(caught.exception))
        self.assertFalse(caught.exception.fatal)
        server = self.serve((503, {}), (503, {}), (503, {}))
        provider = fit.make_provider({**ai, "base_url": server.url}, env={})
        provider.sleep = lambda seconds: None
        with self.assertRaises(FitError) as caught:
            provider.complete("s", "p")
        self.assertIn("busy or having trouble (HTTP 503)", str(caught.exception))
        self.assertFalse(caught.exception.fatal)

    def test_unusual_answers(self):
        cases = [(chat_reply("{}", finish="length"), "cut off"), (chat_reply(None, refusal="no"), "declined"),
                 (chat_reply(""), "empty answer"), ({"id": "x"}, "Chat Completions format"),
                 (b"not json", "Chat Completions format")]
        for body, message in cases:
            server = self.serve((200, body))
            ai, _ = self.ai(base_url=server.url)
            with self.assertRaises(FitError) as caught:
                fit.make_provider(ai, env={}).complete("s", "p")
            self.assertIn(message, str(caught.exception))
        parts = chat_reply([{"type": "text", "text": json.dumps(GOOD_ANSWER)}])
        server = self.serve((200, parts))
        ai, _ = self.ai(base_url=server.url)
        self.assertEqual(fit.parse_fit(fit.make_provider(ai, env={}).complete("s", "p"))["fit"], "possible")

    def test_nothing_listening(self):
        server = self.serve()
        url = server.url
        server.close()
        ai, _ = self.ai(base_url=url)
        with self.assertRaises(FitError) as caught:
            fit.make_provider(ai, env={}).complete("s", "p")
        self.assertTrue(caught.exception.fatal)
        self.assertIn("model server is running", str(caught.exception))


class CoreWithoutAiTests(unittest.TestCase):
    def test_core_imports_without_ai_packages(self):
        code = ("import sys\n"
                "for name in ('anthropic', 'markitdown', 'mcp'): sys.modules[name] = None\n"
                "import usajobs_monitor.cli, usajobs_monitor.runner, usajobs_monitor.fit, usajobs_monitor.tools\n"
                "usajobs_monitor.cli.build_parser().parse_args(['fit', '1', '--preview'])\n")
        done = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True,
                              cwd=Path(__file__).resolve().parent.parent)
        self.assertEqual(done.returncode, 0, done.stderr)


if __name__ == "__main__":
    unittest.main()
