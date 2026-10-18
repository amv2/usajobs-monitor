"""Word briefing, Excel tracker and the ``usajobs_last_run.json`` summary."""

from __future__ import annotations

import csv
import html
import io
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from .jobs import (
    agency_display,
    as_text,
    grade_display,
    job_id,
    job_url,
    location_display,
    occupational_series_display,
    pay_plan_display,
    promotion_potential_display,
    salary_display,
)
from .config import target_label
from .presets import preset_labels
from .reasons import code_of
from .review import SENIOR_EXECUTIVE_PAY_PLANS, ReviewResult, applicant_label, target_agency_hit
from .tracking import CLOSING_SOON_DAYS, RunView, baseline_text, near_miss_reason_text, watch_list_label

WORD_NAME = "usajobs_briefing.docx"
EXCEL_NAME = "usajobs_tracker.xlsx"
HTML_NAME = "usajobs_digest.html"
CSV_NAME = "usajobs_jobs.csv"
SUMMARY_NAME = "usajobs_last_run.json"
REPORT_NAMES = (WORD_NAME, EXCEL_NAME)  # what runs wrote before report.formats existed
FORMAT_FILES = {"docx": WORD_NAME, "xlsx": EXCEL_NAME, "html": HTML_NAME, "csv": CSV_NAME}
REPORT_LABELS = {WORD_NAME: "Word briefing", EXCEL_NAME: "Excel tracker", HTML_NAME: "HTML digest",
                 CSV_NAME: "CSV export"}
QUALIFICATION_CHARS = 2000


def format_timestamp(moment: datetime) -> str:
    """'Refreshed October 7, 2026 at 9:05 AM' without platform-specific strftime flags."""
    hour = moment.hour % 12 or 12
    meridiem = "AM" if moment.hour < 12 else "PM"
    return f"Refreshed {moment:%B} {moment.day}, {moment.year} at {hour}:{moment:%M} {meridiem}"


def resume_note(resumes: list[Path], rated: int = 0) -> str:
    if rated:
        return (f"Resumes: the AI compared your resume with {rated} posting{'s' if rated != 1 else ''}; "
                "see AI fit on each job and the fit notes.")
    if not resumes:
        return "Resumes: none found in the configured folder."
    return f"Resumes: {len(resumes)} on file. They are listed for reference and are not scored against postings."


def resume_pdfs(folder: Path | None) -> list[Path]:
    if folder is None or not folder.is_dir():
        return []
    pdfs = [path for path in folder.rglob("*") if path.is_file() and path.suffix.lower() == ".pdf"]
    return sorted(pdfs, key=lambda p: str(p).casefold())


def _money(value: Any) -> str:
    return f"${float(value or 0):,.0f}"


def _series_list(codes: list[str]) -> str:
    return ", ".join(codes)


def active_filters_text(config: dict[str, Any], decisions_name: str) -> list[str]:
    """Describe the filters actually in effect, one line per rule."""
    filters = config["filters"]
    search = config["search"]
    eligibility = config.get("eligibility") or {}
    lines = []

    salary = float(filters.get("salary_upper_must_exceed") or 0)
    line = (f"Compensation: posted upper salary must exceed {_money(salary)}." if salary
            else "Compensation: no minimum posted salary.")
    if filters.get("min_gs_grade"):
        line += f" GS-scale postings must also reach grade {filters['min_gs_grade']}."
    target = target_label(config)
    if target:
        line += f" Pay floors are set from your target, {target}."
    watch = watch_list_label(config)
    if watch != "Watch list":
        line += f" Postings that fail only this rule are shown on the {watch}."
    lines.append(line)

    routes = ["open to the public"]
    if config.get("profile", {}).get("current_agency_terms") or filters.get("current_agency_terms"):
        routes.append("internal to your current agency")
    if eligibility.get("service") == "excepted":
        routes.append("excepted-service employees")
    if eligibility.get("has_competitive_status"):
        routes.append("competitive status")
    if eligibility.get("interchange_agency"):
        routes.append(f"OPM interchange via {eligibility['interchange_agency']} (explicit instructions or special authorities)")
    extra = eligibility.get("additional_hiring_paths") or []
    if extra:
        routes.append("additional hiring paths: " + ", ".join(extra))
    lines.append(f"Eligibility: a supported route for the {applicant_label(config)} is required ("
                 + "; ".join(routes) + "). Other special eligibilities are not assumed; "
                 "appointment-specific conditions still require verification.")
    if preset_labels(config):
        lines[-1] += " Applicant type: " + "; ".join(preset_labels(config)) + "."

    lines.append("Appointment type: temporary, term, NTE, detail, and other non-permanent appointment types "
                 "(USAJOBS appointment-type code, or the job title) are excluded.")

    plans = sorted(SENIOR_EXECUTIVE_PAY_PLANS) if filters.get("exclude_senior_executive", True) else []
    plans += sorted({str(p).upper() for p in filters.get("excluded_pay_plans", [])} - set(plans))
    if plans:
        line = f"Pay plans excluded: {', '.join(plans)}."
        if filters.get("exclude_senior_executive", True):
            line += " Titles naming senior executive roles are also excluded."
        lines.append(line)

    excluded = filters.get("excluded_job_series") or []
    if excluded:
        lines.append(f"Occupational series: postings listed only under {_series_list(excluded)} are dropped; "
                     "a posting that also lists another series is kept.")

    signals = []
    if search.get("role_title_terms"):
        signals.append("a title term")
    if search.get("target_job_series"):
        signals.append(f"a target job series ({_series_list(search['target_job_series'])})")
    signals.append("a target keyword")
    line = "Relevance: a target role signal (" + ", ".join(signals) + ") is required"
    if search.get("target_agencies"):
        line += "; a target agency raises the ranking but is not enough on its own"
    line += "."
    if filters.get("mismatch_terms") or filters.get("mismatch_job_series"):
        line += " Obvious occupational mismatches are excluded by job title and job series, not announcement boilerplate."
    lines.append(line)

    near_floor = float(filters.get("near_miss_upper_salary_must_exceed") or 0)
    line = "Near misses: postings excluded only for salary, eligibility route, or appointment type"
    if near_floor:
        line += f" whose upper pay exceeds {_money(near_floor)}"
    if filters.get("near_miss_min_gs_grade"):
        line += f" (GS-scale postings must reach grade {filters['near_miss_min_gs_grade']})"
    line += ". Hard-no failures (mismatches, executive roles, excluded series, other agencies' internal postings, no role signal) are never listed."
    lines.append(line)

    lines.append(f"Your decisions: postings marked not interested in {decisions_name} are hidden; "
                 "applied and watching postings are flagged.")
    ai = config.get("ai") or {}
    if ai.get("rank") and ai.get("fit_notes") and ai.get("send_resume") is True:
        lines.append("AI ranking: within each list, postings the AI rated a stronger fit come first, "
                     "then unrated ones, then stretches and unlikely fits. Nothing is hidden, and the AI judges only "
                     "the minimum qualifications, not the assessment questionnaire.")
    return lines


def add_job_paragraph(doc: Any, result: ReviewResult, view: RunView, *, reason: str = "", show_score: bool = False) -> None:
    job = result.job
    p = doc.add_paragraph()
    p.add_run(str(job.get("PositionTitle", "Untitled"))).bold = True
    flags = view.flag_text(result)
    if flags:
        p.add_run(f"  [{flags}]").bold = True
    if show_score:
        p.add_run(f"\nScore: {result.score}")
    p.add_run(f"\nAgency: {agency_display(job)}")
    p.add_run(f"\nPay plan: {pay_plan_display(job)}")
    p.add_run(f"\nGrade: {grade_display(job)}")
    p.add_run(f"\nPromotion potential / full performance level: {promotion_potential_display(job)}")
    p.add_run(f"\nOccupational series: {occupational_series_display(job)}")
    p.add_run(f"\nLocation: {location_display(job) or 'Not listed'}")
    p.add_run(f"\nSalary: {salary_display(job) or 'Not listed'}")
    p.add_run(f"\nClose date: {job.get('ApplicationCloseDate', 'Not listed')}")
    if view.your_status(result):
        p.add_run(f"\nYour status: {view.your_status(result)}")
    if view.fit_text(result):
        summary = view.fit_summaries.get(job_id(job), "")
        p.add_run(f"\nAI fit: {view.fit_text(result)}" + (f". {summary}" if summary else ""))
    if reason:
        p.add_run(f"\nWhy excluded: {reason}")
    p.add_run(f"\nUSAJOBS: {job_url(job)} (control number {job_id(job)})")
    if result.eligibility_basis:
        p.add_run(f"\nEligibility basis: {result.eligibility_basis}")


def write_word(path: Path, results: list[ReviewResult], config: dict[str, Any], resumes: list[Path], calls: int,
               view: RunView, decisions_name: str, now: datetime | None = None) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    now = now or datetime.now()
    watch_label = watch_list_label(config)
    target_agencies = [r for r in view.matches if target_agency_hit(r.job, config["search"].get("target_agencies", []))]
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10)

    title = doc.add_heading(config.get("report", {}).get("title") or "USAJOBS Job Monitor", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph(format_timestamp(now))
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Summary", level=1)
    summary = doc.add_paragraph()
    summary.add_run(f"Jobs reviewed: {len(results)}").bold = True
    summary.add_run(f"\nUSAJOBS API searches run: {calls}")
    summary.add_run(f"\nMatches after final filters: {len(view.matches)}")
    summary.add_run(f"\nTarget-agency matches: {len(target_agencies)}")
    summary.add_run(f"\n{watch_label}: {len(view.watch)}")
    if view.baseline:
        summary.add_run(f"\nNew since last run: baseline run; {baseline_text(view.history_state)}")
    else:
        summary.add_run(f"\nNew since last run: {view.new_count(view.matches)} matches, "
                        f"{view.new_count(view.watch)} watch list, {view.new_count(view.near)} near misses")
    summary.add_run(f"\nClosing within {CLOSING_SOON_DAYS} days: {view.closing_soon_count()}")
    summary.add_run(f"\nDropped since last run: {len(view.dropped)}")
    if view.hidden:
        summary.add_run(f"\nHidden as not interested: {len(view.hidden)}")
    summary.add_run(f"\n{resume_note(resumes, len(view.fits))}")

    if not view.matches:
        doc.add_paragraph(
            "No jobs remained after applying the final filters. The tracker still records reviewed, excluded, "
            "watch-list, and near-miss postings for transparency."
        )
    else:
        doc.add_heading("Top Matches", level=1)
        for result in view.matches[:10]:
            add_job_paragraph(doc, result, view)

    if watch_label != "Watch list" or view.watch:
        doc.add_heading(watch_label[:1].upper() + watch_label[1:].replace("watch list", "Watch List"), level=1)
        doc.add_paragraph(
            "Relevant postings at the watch-list pay plans and grades that passed every check except the salary rule. "
            "They are not matches; they are listed so strong openings at those grades are not missed."
        )
        if not view.watch:
            doc.add_paragraph("No postings qualified for the watch list in this run.")
        for result in view.watch:
            add_job_paragraph(doc, result, view, show_score=True)

    doc.add_heading("Near Misses", level=1)
    doc.add_paragraph(
        "Close calls excluded only for salary, eligibility route, or appointment type. Final filters are unchanged."
    )
    if not view.near:
        doc.add_paragraph("No near-miss postings met the reporting threshold.")
    for result in view.near:
        add_job_paragraph(doc, result, view, reason=near_miss_reason_text(result), show_score=True)

    if view.dropped:
        doc.add_heading("Dropped Since Last Run", level=1)
        for entry in view.dropped:
            label = {"match": "match", "watch": "watch list", "near_miss": "near miss"}.get(entry["list"], entry["list"])
            doc.add_paragraph(f"{entry['title']} ({entry['agency']}): was a {label}; {entry['reason']}. {entry['url']}")

    doc.add_heading("Active Filters", level=1)
    doc.add_paragraph("\n".join(active_filters_text(config, decisions_name)))

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


EXCEL_HEADERS = [
    "Status", "Score", "Title", "Agency", "Department", "Pay Plan", "Grade",
    "Promotion Potential / Full Performance Level", "Occupational Series", "Location", "Salary",
    "Close Date", "USAJOBS URL", "Exclusion Reasons", "Eligibility Basis", "Flags", "Your Status", "Control Number",
    "AI Fit",
]
EXCEL_WIDTHS = [12, 8, 38, 32, 26, 16, 14, 38, 28, 30, 24, 16, 48, 52, 70, 22, 24, 16, 14]


def excel_row(status: str, result: ReviewResult, view: RunView, reasons: str) -> list[Any]:
    job = result.job
    return [
        status, result.score, job.get("PositionTitle", ""), job.get("OrganizationName", ""),
        job.get("DepartmentName", ""), pay_plan_display(job), grade_display(job),
        promotion_potential_display(job), occupational_series_display(job), location_display(job),
        salary_display(job), job.get("ApplicationCloseDate", ""), job_url(job), reasons,
        result.eligibility_basis, view.flag_text(result), view.your_status(result), job_id(job),
        view.fit_text(result),
    ]


def add_sheet(wb: Any, title: str, color: str, rows: list[list[Any]], headers: list[str] = EXCEL_HEADERS,
              widths: list[int] = EXCEL_WIDTHS, first: bool = False) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_excel(path: Path, results: list[ReviewResult], config: dict[str, Any], resumes: list[Path], calls: int,
                view: RunView, decisions_path: Path, now: datetime | None = None) -> None:
    from openpyxl import Workbook

    now = now or datetime.now()
    wb = Workbook()
    match_ids = {id(r) for r in view.matches}
    watch_ids = {id(r) for r in view.watch}
    hidden_ids = {id(r) for r in view.hidden}

    def status(r: ReviewResult) -> str:
        if id(r) in hidden_ids:
            return "Hidden"
        if id(r) in match_ids:
            return "Match"
        return "Watch" if id(r) in watch_ids else "Excluded"

    order = {"Match": 0, "Watch": 1, "Excluded": 2, "Hidden": 3}
    tracker = sorted(results, key=lambda r: (order[status(r)], -r.score, str(r.job.get("PositionTitle", ""))))
    add_sheet(wb, "Tracker", "1F4E78", [excel_row(status(r), r, view, "; ".join(r.reasons)) for r in tracker], first=True)
    add_sheet(wb, "Watch List", "C55A11", [excel_row("Watch", r, view, "; ".join(r.reasons)) for r in view.watch])
    add_sheet(wb, "Near Misses", "8064A2", [excel_row("Near Miss", r, view, near_miss_reason_text(r)) for r in view.near])
    add_sheet(wb, "Dropped", "7F7F7F",
              [[e["title"], e["agency"], e["list"], e["close_date"], e["reason"], e["url"], e["control_number"]] for e in view.dropped],
              headers=["Title", "Agency", "Was On", "Close Date", "Reason", "USAJOBS URL", "Control Number"],
              widths=[38, 32, 14, 16, 18, 48, 16])

    summary = wb.create_sheet("Summary")
    summary_rows = [
        ("Report", config.get("report", {}).get("title") or "USAJOBS Job Monitor"),
        ("Refreshed", now.isoformat(timespec="seconds")),
        ("Jobs reviewed", len(results)),
        ("USAJOBS API searches run", calls),
        ("Matches after final filters", len(view.matches)),
        (watch_list_label(config), len(view.watch)),
        ("Near misses reported", len(view.near)),
        ("New since last run", f"baseline run; {baseline_text(view.history_state)}" if view.baseline else
         f"{view.new_count(view.matches)} matches, {view.new_count(view.watch)} watch list, {view.new_count(view.near)} near misses"),
        ("Dropped since last run", len(view.dropped)),
        ("Hidden as not interested", len(view.hidden)),
        ("Resumes", resume_note(resumes, len(view.fits))),
        ("Salary rule", f"Upper salary must exceed {_money(config['filters'].get('salary_upper_must_exceed'))}"),
        *([("Pay target", target_label(config))] if target_label(config) else []),
        ("Your decisions file", str(decisions_path)),
    ]
    for row in summary_rows:
        summary.append(row)
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 90

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def summary_entry(result: ReviewResult, view: RunView, *, with_reasons: bool = False) -> dict[str, Any]:
    job = result.job
    entry = {
        "title": job.get("PositionTitle", ""),
        "agency": job.get("OrganizationName", ""),
        "url": job_url(job),
        "control_number": job_id(job),
        "score": result.score,
        "salary": salary_display(job),
        "pay_plan": pay_plan_display(job),
        "grade": grade_display(job),
        "promotion_potential": promotion_potential_display(job),
        "close_date": job.get("ApplicationCloseDate", ""),
        "eligibility_basis": result.eligibility_basis,
        "flags": view.flags.get(job_id(job), []),
        "your_status": view.your_status(result),
        "ai_fit": view.fit_text(result),
        "qualification_summary": re.sub(r"\s+", " ", as_text(job.get("QualificationSummary"))).strip()[:QUALIFICATION_CHARS],
    }
    if with_reasons:
        entry["exclusion_reasons"] = [str(r) for r in result.reasons]
        entry["exclusion_codes"] = [code.value if (code := code_of(r)) else "" for r in result.reasons]
    return entry


def report_files(config: dict[str, Any]) -> list[str]:
    """File names for ``report.formats``, in a fixed order."""
    formats = (config.get("report") or {}).get("formats") or ["docx", "xlsx"]
    return [name for fmt, name in FORMAT_FILES.items() if fmt in formats]


def list_status(result: ReviewResult, view: RunView) -> str:
    """Match, Watch, Near Miss, Excluded or Hidden: the list a posting landed on."""
    for label, items in (("Hidden", view.hidden), ("Match", view.matches), ("Watch", view.watch),
                         ("Near Miss", view.near)):
        if any(r is result for r in items):
            return label
    return "Excluded"


def _csv_cell(value: Any) -> Any:
    # Spreadsheet apps run cells that start with these characters as formulas.
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def write_csv(path: Path, results: list[ReviewResult], view: RunView) -> None:
    """Every reviewed posting, one row each, with the list it landed on."""
    order = {"Match": 0, "Watch": 1, "Near Miss": 2, "Excluded": 3, "Hidden": 4}
    rows = sorted(((list_status(r, view), r) for r in results),
                  key=lambda pair: (order[pair[0]], -pair[1].score, str(pair[1].job.get("PositionTitle", ""))))
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\n")
    writer.writerow(["List", *EXCEL_HEADERS[1:]])
    for status, result in rows:
        reasons = near_miss_reason_text(result) if status == "Near Miss" else "; ".join(map(str, result.reasons))
        writer.writerow([_csv_cell(v) for v in excel_row(status, result, view, reasons)])
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(out.getvalue(), encoding="utf-8-sig")


HTML_STYLE = {
    "body": "margin:0;padding:24px 16px;background:#f4f5f7;font-family:-apple-system,'Segoe UI',Roboto,Helvetica,"
            "Arial,sans-serif;color:#1f2328;font-size:15px;line-height:1.45",
    "wrap": "max-width:760px;margin:0 auto",
    "card": "background:#ffffff;border:1px solid #d8dee4;border-radius:8px;padding:14px 16px;margin:0 0 12px",
    "h1": "font-size:22px;margin:0 0 4px",
    "h2": "font-size:17px;margin:24px 0 8px;padding-bottom:4px;border-bottom:2px solid #1f4e78",
    "muted": "color:#59636e;font-size:13px;margin:0",
    "title": "font-size:16px;font-weight:600;margin:0 0 6px",
    "flag": "display:inline-block;background:#fff1c2;color:#7a4a00;border-radius:4px;padding:0 6px;"
            "font-size:12px;font-weight:600;margin-left:6px",
    "fact": "margin:2px 0;font-size:14px",
    "link": "color:#0b5cad",
    "stat": "display:inline-block;margin:0 18px 6px 0",
}


def _safe_link(url: str) -> str:
    return url if re.match(r"https://[A-Za-z0-9.-]+/", url or "") else ""


def _html_job(result: ReviewResult, view: RunView, reason: str = "") -> str:
    job = result.job
    esc = html.escape
    flags = "".join(f'<span style="{HTML_STYLE["flag"]}">{esc(flag)}</span>'
                    for flag in view.flags.get(job_id(job), []))
    facts = [
        ("Agency", agency_display(job)),
        ("Grade", f"{pay_plan_display(job)} {grade_display(job)}".strip()),
        ("Promotion potential", promotion_potential_display(job)),
        ("Series", occupational_series_display(job)),
        ("Location", location_display(job) or "Not listed"),
        ("Salary", salary_display(job) or "Not listed"),
        ("Closes", str(job.get("ApplicationCloseDate", "") or "Not listed").partition("T")[0]),
    ]
    if view.your_status(result):
        facts.append(("Your status", view.your_status(result)))
    if view.fit_text(result):
        summary = view.fit_summaries.get(job_id(job), "")
        facts.append(("AI fit", view.fit_text(result) + (f". {summary}" if summary else "")))
    if reason:
        facts.append(("Why excluded", reason))
    if result.eligibility_basis:
        facts.append(("Eligibility", result.eligibility_basis))
    parts = [f'<div style="{HTML_STYLE["card"]}"><p style="{HTML_STYLE["title"]}">'
             f'{esc(str(job.get("PositionTitle", "Untitled")))}{flags}</p>']
    parts += [f'<p style="{HTML_STYLE["fact"]}"><b>{esc(k)}:</b> {esc(str(v))}</p>' for k, v in facts]
    link = _safe_link(job_url(job))
    number = esc(job_id(job))
    parts.append(f'<p style="{HTML_STYLE["fact"]}">'
                 + (f'<a style="{HTML_STYLE["link"]}" href="{esc(link, quote=True)}">View on USAJOBS</a> · '
                    if link else "")
                 + f'control number {number}</p></div>')
    return "".join(parts)


def write_html(path: Path, results: list[ReviewResult], config: dict[str, Any], calls: int, view: RunView,
               decisions_name: str, now: datetime | None = None) -> None:
    """A self-contained HTML digest of the whole run, readable in any browser or mail client."""
    now = now or datetime.now()
    esc = html.escape
    title = config.get("report", {}).get("title") or "USAJOBS Job Monitor"
    watch_label = watch_list_label(config)
    stats = [("Matches", len(view.matches)), (watch_label, len(view.watch)), ("Near misses", len(view.near)),
             ("Closing soon", view.closing_soon_count()), ("Jobs reviewed", len(results)), ("Searches", calls)]
    body = [f'<div style="{HTML_STYLE["wrap"]}">',
            f'<h1 style="{HTML_STYLE["h1"]}">{esc(title)}</h1>',
            f'<p style="{HTML_STYLE["muted"]}">{esc(format_timestamp(now))}</p>',
            f'<div style="{HTML_STYLE["card"]};margin-top:16px">'
            + "".join(f'<span style="{HTML_STYLE["stat"]}"><b>{n}</b> {esc(label)}</span>' for label, n in stats)
            + "</div>"]
    if view.baseline:
        body.append(f'<p style="{HTML_STYLE["muted"]}">Baseline run: {esc(baseline_text(view.history_state))}.</p>')
    else:
        body.append(f'<p style="{HTML_STYLE["muted"]}">New since last run: {view.new_count(view.matches)} matches, '
                    f"{view.new_count(view.watch)} watch list, {view.new_count(view.near)} near misses. "
                    f"Dropped: {len(view.dropped)}. Hidden as not interested: {len(view.hidden)}.</p>")

    def section(heading: str, note: str, items: list[str], empty: str) -> None:
        body.append(f'<h2 style="{HTML_STYLE["h2"]}">{esc(heading)}</h2>')
        if note:
            body.append(f'<p style="{HTML_STYLE["muted"]};margin-bottom:8px">{esc(note)}</p>')
        body.extend(items or [f'<p style="{HTML_STYLE["muted"]}">{esc(empty)}</p>'])

    section("Matches", "", [_html_job(r, view) for r in view.matches], "No jobs met the final filters in this run.")
    if watch_label != "Watch list" or view.watch:
        section(watch_label[:1].upper() + watch_label[1:], "Relevant postings that failed only the pay rule.",
                [_html_job(r, view) for r in view.watch], "No postings qualified for the watch list.")
    section("Near misses", "Excluded only for salary, eligibility route, or appointment type.",
            [_html_job(r, view, near_miss_reason_text(r)) for r in view.near],
            "No near-miss postings met the reporting threshold.")
    if view.dropped:
        section("Dropped since last run", "", [
            f'<p style="{HTML_STYLE["fact"]}">{esc(e["title"])} ({esc(e["agency"])}): {esc(e["reason"])}</p>'
            for e in view.dropped], "")
    body.append(f'<h2 style="{HTML_STYLE["h2"]}">Active filters</h2>')
    body.extend(f'<p style="{HTML_STYLE["fact"]}">{esc(line)}</p>' for line in active_filters_text(config, decisions_name))
    body.append("</div>")
    page = ('<!doctype html>\n<html lang="en"><head><meta charset="utf-8">'
            '<meta name="viewport" content="width=device-width, initial-scale=1">'
            f"<title>{esc(title)}</title></head>"
            f'<body style="{HTML_STYLE["body"]}">' + "".join(body) + "</body></html>\n")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(page, encoding="utf-8")
