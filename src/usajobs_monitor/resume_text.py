"""Read a resume into plain text, on this computer.

Markdown and text files are read as they are. Word files use python-docx,
which the core already installs. PDF and other formats use the optional
``markitdown`` package from the ``ai`` extra. Nothing here touches the network.
"""

from __future__ import annotations

import os
from pathlib import Path

TEXT_SUFFIXES = {".md", ".markdown", ".txt", ".text"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | {".docx", ".pdf"}
MAX_RESUME_CHARS = 60_000  # about 15 pages; longer files are almost always the wrong file
MAX_RESUMES = 5  # versions compared at once; each one adds to every prompt


class ResumeError(Exception):
    """A fixed, safe message naming the file but never its contents."""


def find_resumes(resume: Path | None, folder: Path | None) -> list[Path]:
    return pick_resumes(resume, folder)[0]


def pick_resumes(resume: Path | None, folder: Path | None) -> tuple[list[Path], int]:
    """``(resumes, left_out)``: ``ai.resume`` when it is a file, else the supported files
    in that folder or ``paths.resume_folder``: versions of one person's resume, and the
    AI judges each posting with the one that fits it best. A folder you keep adding
    tailored copies to can outgrow ``MAX_RESUMES``; then the most recently changed
    ones are used and ``left_out`` counts the rest."""
    if resume is not None:
        if resume.is_dir():
            folder = resume
        elif resume.is_file():
            return [resume], 0
        else:
            raise ResumeError(f"ai.resume: {resume} was not found")
    if folder is None:
        raise ResumeError("set ai.resume to your resume file (.md, .txt, .docx or .pdf)")
    if not folder.is_dir():
        raise ResumeError(f"paths.resume_folder: {folder} was not found")
    candidates = sorted((p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES
                         and not p.name.startswith((".", "~$"))), key=lambda p: p.name.casefold())
    if not candidates:
        raise ResumeError(f"no .md, .txt, .docx or .pdf resume in {folder}; set ai.resume to one file")
    if len(candidates) <= MAX_RESUMES:
        return candidates, 0
    newest = sorted(candidates, key=lambda p: p.stat().st_mtime, reverse=True)[:MAX_RESUMES]
    # Name order, not age order, so the prompt (and its cache) stays the same between runs.
    return sorted(newest, key=lambda p: p.name.casefold()), len(candidates) - MAX_RESUMES


def resume_names(paths: list[Path]) -> list[str]:
    """Short names for several resume versions: file names without the suffix, and
    without the words they all start with ("Pat Lee - Resume - Analyst" and
    "Pat Lee - Resume - Director" become "Analyst" and "Director")."""
    stems = [path.stem.replace('"', "'") for path in paths]
    if len(stems) < 2:
        return stems
    shared = len(os.path.commonprefix(stems))
    # Cut only at a separator, so a word is never split: "Resume v1" and "Resume v2"
    # become "v1" and "v2", not "1" and "2".
    while shared and stems[0][shared - 1] not in " -_–,.":
        shared -= 1
    short = [stem[shared:].strip(" -_–,.") for stem in stems]
    return short if all(short) and len(set(short)) == len(short) else stems


def read_resumes(paths: list[Path]) -> str:
    """One resume as plain text, or several, each marked with its short name."""
    if len(paths) == 1:
        return read_resume(paths[0])
    return "\n\n".join(f'<resume_version name="{name}">\n{read_resume(path)}\n</resume_version>'
                        for name, path in zip(resume_names(paths), paths))


def read_resume(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix in TEXT_SUFFIXES:
            text = path.read_text(encoding="utf-8-sig", errors="replace")
        elif suffix == ".docx":
            text = _docx_text(path)
        else:
            text = _markitdown_text(path)
    except ResumeError:
        raise
    except Exception as exc:  # noqa: BLE001 - parser messages can quote file contents
        raise ResumeError(f"could not read {path.name} ({type(exc).__name__}); save it as .md or .txt and "
                          "point ai.resume at that file") from None
    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if not text:
        raise ResumeError(f"{path.name} has no readable text (a scanned PDF?); save it as .md or .txt")
    if len(text) > MAX_RESUME_CHARS:
        raise ResumeError(f"{path.name} has {len(text):,} characters of text, more than the {MAX_RESUME_CHARS:,} "
                          "a resume should need; check that ai.resume points at your resume")
    return text


def _docx_text(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = list(dict.fromkeys(cell.text.strip() for cell in row.cells if cell.text.strip()))
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _markitdown_text(path: Path) -> str:
    try:
        from markitdown import MarkItDown
    except ImportError:
        raise ResumeError(f"reading {path.suffix or 'this'} files needs the ai extra "
                          '(pip install "usajobs-monitor[ai]"), or save the resume as .md or .txt') from None
    # Local conversion only: no plugins, no LLM image captions, no URL fetching.
    return MarkItDown(enable_plugins=False).convert_local(str(path)).text_content
